"""
Margin — ask your documents.

An end-to-end RAG (Retrieval-Augmented Generation) chatbot built with
Streamlit + LangChain + Google Gemini, following the same pipeline as the
Practical Generative AI Workshop notebook:

    Upload -> Parse -> Chunk -> Embed -> Vector store -> Retrieve -> Generate

Each stage is a small, clearly named function below, so the code reads top to
bottom in the same order as the pipeline.
"""

import hashlib
import html
import io
import os
import re
import time
from pathlib import Path

import numpy as np
import pandas as pd
import streamlit as st
from langchain_core.documents import Document
from langchain_core.messages import AIMessage, HumanMessage, SystemMessage
from langchain_core.vectorstores import InMemoryVectorStore
from langchain_text_splitters import RecursiveCharacterTextSplitter

# ---------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------

SAMPLE_PATH = Path(__file__).parent / "sample_docs" / "lumen_design_handbook.md"
SAMPLE_LABEL = "Sample: Lumen design handbook"
UPLOAD_LABEL = "Upload my own files"

NOT_FOUND = "I couldn't find that in your documents."
MAX_CHUNKS = 400              # keeps free-tier embedding calls reasonable
HISTORY_TURNS = 6             # short-term memory: last N messages resent to the model
CHAT_MODELS = ["gemini-3.1-flash-lite", "gemini-3.5-flash"]
FALLBACK_EMBED_MODEL = "models/gemini-embedding-001"

SAMPLE_STARTERS = [
    "How long does a new component request take?",
    "What contrast ratio does body text need?",
    "When are design crits and how do I book one?",
    "What has to be done before handoff to engineering?",
]
UPLOAD_STARTERS = [
    "Summarise the main topics in these documents.",
    "What are the key numbers or deadlines mentioned?",
    "List any rules or requirements I should know.",
]

SAMPLE_TESTS = [
    {"Question": "How long does a new component take from approval to release?",
     "Answer should contain (any of, comma-separated)": "10 business days, ten business days",
     "Should say it doesn't know": False},
    {"Question": "What contrast ratio is required for body text?",
     "Answer should contain (any of, comma-separated)": "4.5:1",
     "Should say it doesn't know": False},
    {"Question": "On which day do design crits happen?",
     "Answer should contain (any of, comma-separated)": "Wednesday",
     "Should say it doesn't know": False},
    {"Question": "What is the base spacing unit?",
     "Answer should contain (any of, comma-separated)": "4px, 4 px, 4-pixel, 4 pixel",
     "Should say it doesn't know": False},
    {"Question": "How long are deprecated components kept?",
     "Answer should contain (any of, comma-separated)": "two, 2",
     "Should say it doesn't know": False},
    {"Question": "What is the company's parental leave policy?",
     "Answer should contain (any of, comma-separated)": "",
     "Should say it doesn't know": True},
    {"Question": "Who won the last football World Cup?",
     "Answer should contain (any of, comma-separated)": "",
     "Should say it doesn't know": True},
]
UPLOAD_TESTS = [
    {"Question": "What is this document about?",
     "Answer should contain (any of, comma-separated)": "",
     "Should say it doesn't know": False},
    {"Question": "What is the capital of France?",
     "Answer should contain (any of, comma-separated)": "",
     "Should say it doesn't know": True},
]

# ---------------------------------------------------------------------------
# Page setup and styling
# ---------------------------------------------------------------------------

st.set_page_config(page_title="Margin · Ask your documents", page_icon="🖍️", layout="wide")

st.markdown(
    """
<style>
@import url('https://fonts.googleapis.com/css2?family=Figtree:wght@400;500;600&family=Newsreader:opsz,wght@6..72,500;6..72,600&display=swap');
html, body, [class*="st-"], .stMarkdown, button, input, textarea { font-family: 'Figtree', system-ui, sans-serif; }
h1, h2, h3, .margin-title { font-family: 'Newsreader', Georgia, serif !important; letter-spacing: -0.01em; }
.margin-title { font-size: 2.6rem; font-weight: 600; line-height: 1.1; margin: 0; }
.margin-sub { opacity: .72; margin: .35rem 0 1.2rem 0; max-width: 62ch; }
.margin-source { border-left: 3px solid #4A63D9; padding: .55rem .9rem; margin: .5rem 0 .8rem 0;
                 background: rgba(74, 99, 217, .06); border-radius: 0 6px 6px 0; line-height: 1.55; }
.margin-source.cited { border-left-color: #E2B100; background: rgba(255, 214, 0, .10); }
.margin-source .meta { font-size: .82rem; opacity: .75; margin-bottom: .3rem; }
.margin-source mark { background: rgba(255, 214, 0, .45); color: inherit; padding: 0 .1em; border-radius: 2px; }
.margin-pill { display: inline-block; font-size: .78rem; padding: .1rem .55rem; border-radius: 999px;
               border: 1px solid rgba(128,128,128,.35); margin-right: .35rem; }
.margin-step { font-family: 'Newsreader', Georgia, serif; font-size: 1.25rem; font-weight: 600; margin: 1.1rem 0 .2rem 0; }
</style>
""",
    unsafe_allow_html=True,
)


# ---------------------------------------------------------------------------
# Secrets: API key and optional password
# ---------------------------------------------------------------------------

def read_secret(name):
    try:
        value = st.secrets.get(name)
    except Exception:
        value = None
    return value or os.environ.get(name)


def password_gate():
    """If APP_PASSWORD is set in secrets, ask for it before showing the app."""
    password = read_secret("APP_PASSWORD")
    if not password or st.session_state.get("unlocked"):
        return
    st.markdown('<p class="margin-title">Margin</p>', unsafe_allow_html=True)
    entered = st.text_input("Enter the access password to continue", type="password")
    if entered:
        if entered == password:
            st.session_state.unlocked = True
            st.rerun()
        st.error("That password is incorrect. Ask the app owner for the current one.")
    st.stop()


password_gate()
API_KEY = read_secret("GOOGLE_API_KEY")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def get_text(response):
    """Plain text of an LLM response, whether .content is a string or a list of blocks
    (same helper as in the workshop notebook)."""
    content = response.content
    if isinstance(content, str):
        return content
    return "".join(
        block.get("text", "") for block in content
        if isinstance(block, dict) and block.get("type") == "text"
    )


def friendly_error(err):
    msg = str(err)
    if "429" in msg or "RESOURCE_EXHAUSTED" in msg or "quota" in msg.lower():
        return "The Gemini free-tier limit was reached. Wait a minute and try again."
    if "API key" in msg or "API_KEY" in msg or "401" in msg or "403" in msg:
        return "Gemini rejected the API key. Check GOOGLE_API_KEY in the app's secrets."
    if "404" in msg or "not found" in msg.lower():
        return "That Gemini model isn't available for this key. Pick another model in Settings."
    return f"Something went wrong while talking to Gemini: {msg[:300]}"


def source_label(meta):
    parts = [meta.get("source", "document")]
    if meta.get("page"):
        parts.append(f"page {meta['page']}")
    if meta.get("section"):
        parts.append(meta["section"])
    return " · ".join(parts)


STOPWORDS = set("""a an and are as at be by can do does for from has have how i in is it its me my of on or
should that the their this to was what when where which who why will with you your""".split())


def highlight(text, query):
    """Escape text and wrap the query's meaningful words in <mark> so people can see why a chunk matched."""
    safe = html.escape(text)
    words = {w for w in re.findall(r"[A-Za-z0-9.:]+", query.lower()) if len(w) > 2 and w not in STOPWORDS}
    for w in sorted(words, key=len, reverse=True):
        safe = re.sub(rf"(?i)\b({re.escape(w)})", r"<mark>\1</mark>", safe)
    return safe


# ---------------------------------------------------------------------------
# Models (cached so they're created once per API key)
# ---------------------------------------------------------------------------

@st.cache_resource(show_spinner=False)
def get_embeddings(api_key):
    """Pick an embedding model the key can use, like the notebook does, and wrap it in LangChain."""
    from langchain_google_genai import GoogleGenerativeAIEmbeddings

    model_name = FALLBACK_EMBED_MODEL
    try:
        from google import genai
        client = genai.Client(api_key=api_key)
        names = [m.name for m in client.models.list()
                 if "embedContent" in (m.supported_actions or [])]
        if names:
            preferred = [n for n in names if "gemini-embedding" in n]
            model_name = (preferred or names)[0]
    except Exception:
        pass
    return GoogleGenerativeAIEmbeddings(model=model_name, google_api_key=api_key), model_name


def get_llm(api_key, model, temperature):
    from langchain_google_genai import ChatGoogleGenerativeAI
    return ChatGoogleGenerativeAI(model=model, temperature=temperature,
                                  max_output_tokens=1024, google_api_key=api_key)


# ---------------------------------------------------------------------------
# Pipeline stage 1-2: upload and parse
# ---------------------------------------------------------------------------

def clean_text(text):
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_markdown_sections(text, name):
    """Structure-aware parsing: keep each markdown heading with its own section."""
    docs, title, buf = [], None, []
    for line in text.splitlines():
        if line.lstrip().startswith("#"):
            if "\n".join(buf).strip():
                docs.append(Document(page_content=clean_text("\n".join(buf)),
                                     metadata={"source": name, "section": title}))
            title, buf = line.lstrip("# ").strip(), [line.lstrip("# ").strip()]
        else:
            buf.append(line)
    if "\n".join(buf).strip():
        docs.append(Document(page_content=clean_text("\n".join(buf)),
                             metadata={"source": name, "section": title}))
    return docs


def parse_file(name, data):
    """Turn one uploaded file into LangChain Documents (one per page or section)."""
    ext = name.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return [Document(page_content=clean_text(page.extract_text() or ""),
                         metadata={"source": name, "page": i + 1})
                for i, page in enumerate(reader.pages) if (page.extract_text() or "").strip()]
    if ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        lines = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return [Document(page_content=clean_text("\n".join(lines)), metadata={"source": name})]
    text = data.decode("utf-8", errors="ignore")
    if ext == "md":
        return split_markdown_sections(text, name)
    return [Document(page_content=clean_text(text), metadata={"source": name})]


# ---------------------------------------------------------------------------
# Pipeline stage 3-5: chunk, embed, store  (cached per set of files + settings)
# ---------------------------------------------------------------------------

@st.cache_resource(show_spinner=False, max_entries=6)
def build_knowledge_base(files, chunk_size, chunk_overlap, api_key):
    """files: tuple of (filename, bytes). Returns the vector store plus stats for the inspector."""
    timings = {}

    t = time.time()
    pages = []
    for name, data in files:
        pages.extend(parse_file(name, data))
    pages = [p for p in pages if p.page_content]
    timings["parse"] = time.time() - t
    if not pages:
        raise ValueError("No readable text was found. Scanned PDFs (images of text) aren't supported yet.")

    t = time.time()
    splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    chunks = splitter.split_documents(pages)
    truncated = len(chunks) > MAX_CHUNKS
    chunks = chunks[:MAX_CHUNKS]
    for i, c in enumerate(chunks):
        c.metadata["chunk_id"] = i
    timings["chunk"] = time.time() - t

    t = time.time()
    embeddings, embed_model = get_embeddings(api_key)
    store = InMemoryVectorStore(embeddings)
    store.add_documents(chunks)
    timings["embed"] = time.time() - t

    vectors = [v["vector"] for v in store.store.values()]
    return {
        "store": store,
        "chunks": chunks,
        "files": [name for name, _ in files],
        "pages": len(pages),
        "characters": sum(len(p.page_content) for p in pages),
        "embed_model": embed_model,
        "dimensions": len(vectors[0]) if vectors else 0,
        "vectors": vectors,
        "timings": timings,
        "truncated": truncated,
    }


# ---------------------------------------------------------------------------
# Pipeline stage 6: retrieve
# ---------------------------------------------------------------------------

CONDENSE_PROMPT = """You rewrite follow-up questions so they can be searched on their own.
Given the conversation and the latest message, write one standalone search query that
includes any names or topics the message refers to. Return only the query, nothing else.

Conversation:
{history}

Latest message: {question}

Standalone query:"""


def rewrite_query(llm, question, history):
    """Turn 'and how long does that take?' into 'how long does a component request take?'."""
    if not history:
        return question
    transcript = "\n".join(f"{m['role']}: {m['content'][:400]}" for m in history[-HISTORY_TURNS:])
    try:
        rewritten = get_text(llm.invoke(CONDENSE_PROMPT.format(history=transcript, question=question))).strip()
        return rewritten.strip('"') or question
    except Exception:
        return question


def retrieve(kb, query, k, use_mmr):
    """Embed the query, then find the k closest chunks. MMR trades a little similarity for variety."""
    embeddings, _ = get_embeddings(API_KEY)
    query_vector = embeddings.embed_query(query)
    fetch_k = max(k * 3, 10)
    scored = kb["store"].similarity_search_with_score_by_vector(query_vector, k=fetch_k)
    if not use_mmr:
        return scored[:k]
    score_by_id = {d.metadata["chunk_id"]: s for d, s in scored}
    diverse = kb["store"].max_marginal_relevance_search_by_vector(
        query_vector, k=k, fetch_k=fetch_k, lambda_mult=0.6)
    return [(d, score_by_id.get(d.metadata["chunk_id"], 0.0)) for d in diverse]


# ---------------------------------------------------------------------------
# Pipeline stage 7: generate
# ---------------------------------------------------------------------------

SYSTEM_PROMPT = """You are Margin, an assistant that answers questions using ONLY the numbered sources
below, which come from the user's own documents.

Rules:
- Use only facts stated in the sources. Do not use outside knowledge and do not guess.
- Cite the source number in square brackets after each fact, like [1] or [2][3].
- If the sources don't contain the answer, reply exactly: "{not_found}" Then add one short
  sentence suggesting a related question the documents can answer.
- Greetings, and questions about earlier messages in this conversation, can be answered from
  the conversation itself.
- Lead with the direct answer in one or two sentences. Add short bullet points only if they help.

SOURCES:
{sources}"""


def build_messages(question, results, history):
    sources = "\n\n".join(
        f"[{i}] ({source_label(d.metadata)})\n{d.page_content}" for i, (d, _) in enumerate(results, 1))
    messages = [SystemMessage(content=SYSTEM_PROMPT.format(not_found=NOT_FOUND, sources=sources))]
    for m in history[-HISTORY_TURNS:]:
        cls = HumanMessage if m["role"] == "user" else AIMessage
        messages.append(cls(content=m["content"]))
    messages.append(HumanMessage(content=question))
    return messages


def stream_answer(llm, messages):
    for chunk in llm.stream(messages):
        text = get_text(chunk)
        if text:
            yield text


def answer_once(question, kb, settings):
    """Non-streaming version of the full pipeline, used by the Test tab."""
    llm = get_llm(API_KEY, settings["model"], settings["temperature"])
    results = retrieve(kb, question, settings["k"], settings["mmr"])
    answer = get_text(llm.invoke(build_messages(question, results, [])))
    return answer, results


# ---------------------------------------------------------------------------
# Rendering helpers
# ---------------------------------------------------------------------------

def render_sources(sources, query, answer):
    cited = {int(n) for n in re.findall(r"\[(\d+)\]", answer or "")}
    n_cited = len([s for s in sources if s["n"] in cited])
    label = f"Sources · {n_cited} cited of {len(sources)} retrieved" if sources else "Sources"
    with st.expander(label, expanded=False):
        for s in sources:
            is_cited = s["n"] in cited
            st.markdown(
                f'<div class="margin-source{" cited" if is_cited else ""}">'
                f'<div class="meta"><b>[{s["n"]}]</b> {html.escape(s["label"])} · match {s["score"]:.2f}'
                f'{" · used in answer" if is_cited else ""}</div>'
                f'{highlight(s["text"], query)}</div>',
                unsafe_allow_html=True,
            )


def to_source_dicts(results):
    return [{"n": i, "label": source_label(d.metadata), "score": float(s),
             "text": d.page_content, "chunk_id": d.metadata.get("chunk_id")}
            for i, (d, s) in enumerate(results, 1)]


# ---------------------------------------------------------------------------
# Sidebar: knowledge base + settings
# ---------------------------------------------------------------------------

with st.sidebar:
    st.markdown("### Your documents")
    mode = st.radio("Knowledge source", [SAMPLE_LABEL, UPLOAD_LABEL], label_visibility="collapsed")
    uploads = []
    if mode == UPLOAD_LABEL:
        uploads = st.file_uploader("PDF, Word, Markdown or text files", type=["pdf", "docx", "txt", "md"],
                                   accept_multiple_files=True)
    st.caption("Files stay in this browser session and are sent to Google Gemini to be processed. "
               "Don't upload confidential or client data.")

    with st.expander("Answer settings"):
        model = st.selectbox("Chat model", CHAT_MODELS + ["Other…"],
                             help="Flash-lite is faster and cheaper. Flash gives fuller answers.")
        if model == "Other…":
            model = st.text_input("Model name", value="gemini-3.5-flash")
        temperature = st.slider("Temperature", 0.0, 1.0, 0.2, 0.1,
                                help="Low keeps answers consistent and close to the source.")
        k = st.slider("Sources per answer (top-k)", 1, 8, 4)
        use_mmr = st.toggle("Prefer varied sources (MMR)", value=False,
                            help="Avoids retrieving several near-identical chunks.")
        use_rewrite = st.toggle("Understand follow-up questions", value=True,
                                help="Rewrites 'what about that?' into a full question before searching.")

    with st.expander("Chunking settings"):
        chunk_size = st.slider("Chunk size (characters)", 200, 2000, 800, 100)
        chunk_overlap = st.slider("Overlap (characters)", 0, 400, 120, 20)

    if st.button("Clear conversation", width="stretch"):
        st.session_state.messages = []
        st.session_state.last_trace = None
        st.rerun()

settings = {"model": model, "temperature": temperature, "k": k, "mmr": use_mmr}

# ---------------------------------------------------------------------------
# Header + guards
# ---------------------------------------------------------------------------

st.markdown('<p class="margin-title">Margin</p>', unsafe_allow_html=True)
st.markdown('<p class="margin-sub">Ask questions about your documents. Every answer points to the '
            'passage it came from, and says so when the answer isn\'t there.</p>', unsafe_allow_html=True)

if not API_KEY:
    st.error("Add your Gemini API key to start. In Streamlit Cloud open the app's Settings, then Secrets, "
             'and add the line GOOGLE_API_KEY = "your-key". Get a free key at aistudio.google.com.')
    st.stop()

if mode == SAMPLE_LABEL:
    files = ((SAMPLE_PATH.name, SAMPLE_PATH.read_bytes()),)
elif uploads:
    files = tuple((f.name, f.getvalue()) for f in uploads)
else:
    st.info("Upload one or more files in the sidebar to build your knowledge base, "
            "or switch to the sample handbook to try it first.")
    st.stop()

files_key = hashlib.md5(b"".join(n.encode() + d for n, d in files)).hexdigest()
if st.session_state.get("files_key") != files_key:
    st.session_state.files_key = files_key
    st.session_state.messages = []
    st.session_state.last_trace = None

try:
    with st.spinner("Reading, chunking and embedding your documents…"):
        kb = build_knowledge_base(files, chunk_size, chunk_overlap, API_KEY)
except Exception as e:
    st.error(str(e) if isinstance(e, ValueError) else friendly_error(e))
    st.stop()

st.markdown(
    f'<span class="margin-pill">{len(kb["files"])} file{"s" if len(kb["files"]) != 1 else ""}</span>'
    f'<span class="margin-pill">{len(kb["chunks"])} chunks</span>'
    f'<span class="margin-pill">{kb["dimensions"]}-number vectors</span>',
    unsafe_allow_html=True,
)
if kb["truncated"]:
    st.warning(f"Only the first {MAX_CHUNKS} chunks were indexed to stay within free-tier limits.")

st.session_state.setdefault("messages", [])
st.session_state.setdefault("last_trace", None)

tab_chat, tab_how, tab_test = st.tabs(["Chat", "How it works", "Test it"])

# ---------------------------------------------------------------------------
# Tab 1: Chat
# ---------------------------------------------------------------------------

with tab_chat:
    starters = SAMPLE_STARTERS if mode == SAMPLE_LABEL else UPLOAD_STARTERS
    if not st.session_state.messages:
        st.caption("Try one of these to start")
        cols = st.columns(len(starters))
        for col, q in zip(cols, starters):
            if col.button(q, width="stretch"):
                st.session_state.pending = q
                st.rerun()

    for m in st.session_state.messages:
        with st.chat_message(m["role"]):
            st.markdown(m["content"])
            if m["role"] == "assistant" and m.get("sources"):
                render_sources(m["sources"], m.get("query", ""), m["content"])

    question = st.chat_input("Ask about your documents")
    question = question or st.session_state.pop("pending", None)

    if question:
        history = list(st.session_state.messages)
        with st.chat_message("user"):
            st.markdown(question)
        with st.chat_message("assistant"):
            try:
                llm = get_llm(API_KEY, settings["model"], settings["temperature"])
                t0 = time.time()
                search_query = rewrite_query(llm, question, history) if use_rewrite else question
                t1 = time.time()
                results = retrieve(kb, search_query, settings["k"], settings["mmr"])
                t2 = time.time()
                messages = build_messages(question, results, history)
                answer = st.write_stream(stream_answer(llm, messages))
                t3 = time.time()
            except Exception as e:
                st.error(friendly_error(e))
                st.stop()
            sources = to_source_dicts(results)
            render_sources(sources, search_query, answer)

        st.session_state.messages.append({"role": "user", "content": question})
        st.session_state.messages.append({"role": "assistant", "content": answer,
                                          "sources": sources, "query": search_query})
        st.session_state.last_trace = {
            "question": question, "search_query": search_query, "sources": sources,
            "prompt": "\n\n".join(f"[{m.type}]\n{m.content}" for m in messages),
            "answer": answer, "model": settings["model"], "temperature": settings["temperature"],
            "times": {"rewrite": t1 - t0, "retrieve": t2 - t1, "generate": t3 - t2},
        }

# ---------------------------------------------------------------------------
# Tab 2: How it works (pipeline inspector)
# ---------------------------------------------------------------------------

with tab_how:
    st.caption("Each step of the RAG pipeline, using the documents loaded right now.")

    st.markdown('<div class="margin-step">1. Upload and parse</div>', unsafe_allow_html=True)
    st.write(f"Read **{len(kb['files'])}** file(s) into **{kb['pages']}** pages or sections, "
             f"**{kb['characters']:,}** characters of text in {kb['timings']['parse']:.2f}s.")
    st.caption("PDFs are read page by page so answers can cite the page. Markdown is split by headings.")

    st.markdown('<div class="margin-step">2. Chunk</div>', unsafe_allow_html=True)
    sizes = [len(c.page_content) for c in kb["chunks"]]
    st.write(f"Split into **{len(kb['chunks'])}** overlapping chunks of up to {chunk_size} characters "
             f"(average {int(np.mean(sizes))}, overlap {chunk_overlap}).")
    chunk_ix = st.number_input("Look at chunk number", 1, len(kb["chunks"]), 1) - 1
    chosen = kb["chunks"][chunk_ix]
    st.markdown(f'<div class="margin-source"><div class="meta">{html.escape(source_label(chosen.metadata))}'
                f'</div>{html.escape(chosen.page_content)}</div>', unsafe_allow_html=True)

    st.markdown('<div class="margin-step">3. Embed</div>', unsafe_allow_html=True)
    vec = kb["vectors"][chunk_ix]
    st.write(f"Each chunk became a list of **{kb['dimensions']}** numbers using `{kb['embed_model']}` "
             f"({kb['timings']['embed']:.1f}s for all chunks). The first eight for this chunk:")
    st.code(", ".join(f"{x:.4f}" for x in vec[:8]) + ", …")
    st.caption("Chunks with similar meaning end up with similar numbers, even when they use different words.")

    st.markdown('<div class="margin-step">4. Store</div>', unsafe_allow_html=True)
    st.write(f"All **{len(kb['vectors'])}** vectors live in LangChain's `InMemoryVectorStore` for this session.")

    st.markdown('<div class="margin-step">5. Retrieve</div>', unsafe_allow_html=True)
    trace = st.session_state.last_trace
    if not trace:
        st.info("Ask a question in the Chat tab and its retrieval and prompt will appear here.")
    else:
        st.write(f"Question: **{trace['question']}**")
        if trace["search_query"] != trace["question"]:
            st.write(f"Rewritten for search: **{trace['search_query']}**")
        df = pd.DataFrame([{"#": s["n"], "Source": s["label"], "Match": s["score"]} for s in trace["sources"]])
        st.dataframe(df, hide_index=True, width="stretch", column_config={
            "Match": st.column_config.ProgressColumn("Match (cosine similarity)", min_value=0.0,
                                                     max_value=1.0, format="%.2f")})
        st.caption(f"Search took {trace['times']['retrieve']:.2f}s"
                   + (f", rewriting took {trace['times']['rewrite']:.2f}s." if trace["times"]["rewrite"] > 0.01 else "."))

        st.markdown('<div class="margin-step">6. Generate</div>', unsafe_allow_html=True)
        st.write(f"`{trace['model']}` at temperature {trace['temperature']} wrote the answer in "
                 f"{trace['times']['generate']:.1f}s, using only the sources above.")
        with st.expander("See the exact prompt sent to Gemini"):
            st.code(trace["prompt"], language=None)

# ---------------------------------------------------------------------------
# Tab 3: Test it (a tiny eval suite)
# ---------------------------------------------------------------------------

with tab_test:
    st.caption("Check the chatbot against questions with known answers. Edit, add or remove rows, then run.")
    defaults = SAMPLE_TESTS if mode == SAMPLE_LABEL else UPLOAD_TESTS
    cases = st.data_editor(pd.DataFrame(defaults), num_rows="dynamic", width="stretch",
                           key=f"tests_{mode}", hide_index=True)

    if st.button("Run tests", type="primary"):
        rows, progress = [], st.progress(0.0, text="Running tests…")
        for i, row in cases.reset_index(drop=True).iterrows():
            q = str(row.get("Question") or "").strip()
            if not q:
                continue
            expected = [e.strip().lower() for e in str(row.get("Answer should contain (any of, comma-separated)") or "").split(",") if e.strip()]
            should_refuse = bool(row.get("Should say it doesn't know"))
            try:
                t = time.time()
                answer, results = answer_once(q, kb, settings)
                latency = time.time() - t
            except Exception as e:
                rows.append({"Result": "⚠️ Error", "Question": q, "Answer": friendly_error(e)})
                continue
            low = answer.lower()
            refused = "couldn't find" in low or "could not find" in low
            cited = bool(re.search(r"\[\d+\]", answer))
            if should_refuse:
                ok, why = refused, "Correctly said it doesn't know" if refused else "Answered when it should have refused"
            else:
                has_expected = (not expected) or any(e in low for e in expected)
                ok = has_expected and not refused and cited
                why = ("Correct and cited" if ok else
                       "Said it doesn't know" if refused else
                       "Missing expected answer" if not has_expected else "No citation")
            rows.append({"Result": "✅ Pass" if ok else "❌ Fail", "Question": q, "Why": why,
                         "Top source": source_label(results[0][0].metadata) if results else "",
                         "Seconds": round(latency, 1), "Answer": answer})
            progress.progress((i + 1) / len(cases), text=f"Ran {i + 1} of {len(cases)}")
        progress.empty()
        if rows:
            passed = sum(r["Result"].startswith("✅") for r in rows)
            st.metric("Passed", f"{passed} of {len(rows)}")
            st.dataframe(pd.DataFrame(rows), hide_index=True, width="stretch")
            st.caption("A failing test is useful: try a different chunk size, top-k or model, then run again.")
